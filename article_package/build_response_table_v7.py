# -*- coding: utf-8 -*-
"""Full Ukrainian response table for Article 1 v7 revision."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Таблиця_відповідей_рецензентам_стаття1_v7c.docx"
FONT = "Times New Roman"
MS_VERSION = "Стаття_Аспірант_Синюк_HAIT_aligned_final_v7c.docx"

ROWS = [
    (
        "1",
        "1",
        "Джерела 23–30 є в списку літератури, але не цитуються в тексті.",
        "Враховано. Джерела [23]–[31] процитовано в огляді літератури/обговоренні; список оновлено сучасними роботами (див. також п. рец. 2 щодо віку джерел).",
        "Literature review; References",
    ),
    (
        "1",
        "2",
        "Не визначено RC та індекс плавності.",
        "Враховано. Додано означення SI = −log₁₀(mean(j_k²)); CV = (SD/|mean|)×100%; RC_disp = 1.96·√2·SD як міжсесійний коефіцієнт дисперсії (не classical repeatability).",
        "Materials and Methods; Table 3 Note",
    ),
    (
        "1",
        "3",
        "Згадано кути плечей/стегон, X-factor тощо, але не наведено в результатах.",
        "Враховано. Уточнено, що ці метрики експортуються покадрово, але зарезервовані для майбутнього аналізу і не входять до session-level результатів цієї статті.",
        "Landmark extraction",
    ),
    (
        "1",
        "4",
        "Розбите речення у Conclusions.",
        "Враховано. Conclusions подано суцільними абзацами без розриву речення.",
        "Conclusions",
    ),
    (
        "1",
        "5",
        "Ідентичні рядки max speed / downswing peak speed у таблиці.",
        "Враховано. У примітці до Table 3 зазначено, що збіг очікуваний (пік швидкості під час даунсвінгу).",
        "Table 3 Note",
    ),
    (
        "1*",
        "6",
        "Жирні рубрики в Abstract EN/UK.",
        "Враховано. EN/UK анотації зі структурними рубриками жирним; EN ≈305 слів (300–350).",
        "ABSTRACT / АНОТАЦІЯ",
    ),
    (
        "1*",
        "7",
        "Збільшити рисунки; читабельність у 2 колонках.",
        "Враховано. Fig. 1, 3–7 — full-width islands (~16.5 см, без розтягування); Fig. 2 — у колонці. Двоколонковий макет HAIT збережено.",
        "Fig. 1–7",
    ),
    (
        "1*",
        "8",
        "DOI / Financing / Accessibility / AI.",
        "Враховано. DOI рукопису: https://doi.org/10.15276/hait.09.2026.5 (також у DOI-рядку першої сторінки). Funding: дослідження без зовнішнього фінансування. Data availability: агрегати в статті + публічний репозиторій https://github.com/ivansinuyk/vr_motion; сирі відео — за запитом. AI: Claude Fable, GPT Astra, GPT-5.6 Sol/Composer у Cursor (2026). Conflicts of Interest зазначено.",
        "DOI line; declarations after REFERENCES",
    ),
    (
        "1*",
        "9",
        "Фото авторів (Oleksii / Karthik).",
        "Враховано. Вбудовано фото Oleksii M. Maksymov та Karthik Iyer у таблицю About the Authors (без перестворення таблиці).",
        "ABOUT THE AUTHORS",
    ),
    (
        "2",
        "1",
        "Змішано реалізоване/заплановане; немає параметрів, детектора, provenance 71 сесій.",
        "Враховано. Table 1a (scientific profile); MediaPipe 1.0.1; ~15 спортсменів без ID у експорті; residual-adaptive R реалізовано, confidence-weighted R_k — optional; MoCap/paired validation — planned.",
        "Table 1; Table 1a; Methods",
    ),
    (
        "2",
        "2",
        "Помилка розмірності scale/speed; нечитабельні формули; немає SI.",
        "Враховано. Формули відновлено; s_k [м/піксель], v_k = (||Δp||/Δt)·s_k (множення, як у коді); CV/RC_disp/SI визначено.",
        "Dynamic scale; kinematics",
    ),
    (
        "2",
        "3",
        "Динамічне масштабування / foreshortening; не називати метричними GT.",
        "Враховано. Показники названо видимими 2D-оцінками площини зображення за L_ref=1.0 м; без каліброваної 3D/багатокамерної моделі.",
        "Dynamic scale; Discussion; Abstract",
    ),
    (
        "2",
        "4",
        "«Repeatability» не є test–retest; перейменувати на between-session dispersion.",
        "Враховано. Розділ/таблиця/рис. 3/keywords/висновки перейменовано; «least dispersed in this corpus»; RC_disp без претензії на classical repeatability.",
        "Results; Table 3; Fig. 3",
    ),
    (
        "2",
        "5",
        "Валідація/абляція не доводять ефективність; Table 7 аномальна; потрібна paired ablation.",
        "Враховано. Keyframe errors залишено як діагностику. Table 7 повністю перераховано paired ablation на тих самих 71 сесіях (спільні похідні й time base): median_only / Kalman / Kalman+RTS / full pipeline; інтерпретація оновлена (відхилення↓ у full pipeline; jerk↓ на RTS; підвищення jerk після blend до raw пояснено).",
        "Table 7; Results; Discussion",
    ),
    (
        "2",
        "6",
        "Анотація 300–350; Fig.1/4; табл.; Data Availability; сучасна література (>50% не старші 10 років).",
        "Враховано. Abstract структуровано; figures/tables islands; Data availability з GitHub; замінено 10 застарілих вторинних джерел на роботи 2017–2022; частка джерел старше 10 років знижена до ≈22.6% (класичні Kalman/RTS/SG/Bland/Winter збережено як першоджерела методу).",
        "ABSTRACT; figures; tables; References",
    ),
    (
        "2",
        "7",
        "Tables 3 і 7 надто вузькі в 2-колонковому макеті.",
        "Враховано. Table 1a, Table 3 і Table 7 винесено у full-width continuous section islands (як і щільні рисунки).",
        "Table 1a; Table 3; Table 7",
    ),
]


def set_run_font(run, size=10, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)


def set_cell(cell, text, bold=False, size=9):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.5)
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("ТАБЛИЦЯ")
    set_run_font(r, 14, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Відповіді авторів на зауваження рецензентів\n"
        "«Markerless video-based golf-stick motion analysis using Kalman filtering "
        "and Rauch-Tung-Striebel smoothing»\n"
        f"Виправлена версія рукопису: {MS_VERSION}"
    )
    set_run_font(r, 11, bold=True)

    note = doc.add_paragraph()
    r = note.add_run(
        "Примітка. «1*» — редакційні помітки з анотованої копії. Усі пункти нижче "
        "відображають стан після ревізії v7c: змістовні правки рецензентів (фото авторів, "
        "table/figure islands, paired ablation re-run, оновлення літератури, DOI/declarations) "
        "плюс синхронізація HAIT-оформлення (шрифти abstract/keywords/references, "
        "підписи таблиць/рисунків, Source-рядки)."
    )
    set_run_font(r, 9)

    table = doc.add_table(rows=1 + len(ROWS), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "№ рецензента",
        "Зауваження рецензента",
        "Відповідь авторів",
        "Місце змін",
    ]
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=9)
    for ri, (rev, num, q, a, page) in enumerate(ROWS, start=1):
        set_cell(table.rows[ri].cells[0], f"{rev}\nп.{num}", bold=True, size=9)
        set_cell(table.rows[ri].cells[1], q, size=9)
        set_cell(table.rows[ri].cells[2], a, size=9)
        set_cell(table.rows[ri].cells[3], page, size=9)
    widths = [Cm(2.0), Cm(5.2), Cm(7.2), Cm(2.8)]
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    for path in (OUT,):
        try:
            doc.save(str(path))
            print("Wrote", path.name)
        except PermissionError:
            alt = path.with_name(path.stem + "_alt.docx")
            doc.save(str(alt))
            print("locked;", path.name, "-> wrote", alt.name)


if __name__ == "__main__":
    main()
